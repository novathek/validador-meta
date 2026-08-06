import openpyxl
import json
from datetime import datetime, date

def to_str(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    return s if s else None

def normalize_header(h):
    if not h:
        return None
    h = str(h).strip()
    replacements = {
        '\u00e1': 'A', '\u00e9': 'E', '\u00ed': 'I', '\u00f3': 'O', '\u00fa': 'U',
        '\u00c1': 'A', '\u00c9': 'E', '\u00cd': 'I', '\u00d3': 'O', '\u00da': 'U',
        '\u00f1': 'N', '\u00d1': 'N',
    }
    for k, v in replacements.items():
        h = h.replace(k, v)
    h = h.upper().replace(' ', '_')
    return h

def inspect_xlsx(filename):
    wb = openpyxl.load_workbook(filename)
    result = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        raw_headers = [cell.value for cell in ws[1]]
        norm_headers = [normalize_header(h) for h in raw_headers]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            obj = {}
            for i, h in enumerate(norm_headers):
                if h and i < len(row):
                    obj[h] = to_str(row[i])
            if any(v is not None for v in obj.values()):
                rows.append(obj)
        result[sheet_name] = {
            'raw_headers': raw_headers,
            'norm_headers': norm_headers,
            'count': len(rows),
            'sample': rows[:1] if rows else []
        }
    return result

for fname in ['directivos.xlsx', 'supervisores.xlsx', 'docentes.xlsx']:
    print(f'\n=== {fname} ===')
    info = inspect_xlsx(fname)
    for sheet, data in info.items():
        print(f'  Hoja: {sheet}')
        print(f'  Columnas originales: {data["raw_headers"]}')
        print(f'  Columnas normalizadas: {data["norm_headers"]}')
        print(f'  Total filas: {data["count"]}')
        if data['sample']:
            print(f'  Muestra fila 1: {json.dumps(data["sample"][0], ensure_ascii=False)}')
